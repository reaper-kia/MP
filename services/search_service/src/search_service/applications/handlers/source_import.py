"""Импорт источников из XLSX (B2-10): колонка A — ссылка, B — название, C — «региональный»."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from search_service.domain.errors import SearchInvalidParams

MAX_IMPORT_ROWS = 500

_YES = {"да", "д", "yes", "y", "true", "1", "+"}
_NO = {"нет", "н", "no", "n", "false", "0", "-"}


@dataclass(frozen=True, slots=True)
class ImportRow:
    row: int  # номер строки в Excel, с 1
    url: str
    name: str | None
    is_regional: bool


@dataclass(frozen=True, slots=True)
class BadRow:
    row: int
    value: str
    message: str


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _looks_like_link(value: str) -> bool:
    return "." in value or "/" in value


def parse_regional(value: str) -> bool | None:
    """«да»/«нет» → bool, пусто → True (как у CreateSource); непонятное → None."""
    lowered = value.strip().lower()
    if not lowered:
        return True
    if lowered in _YES:
        return True
    if lowered in _NO:
        return False
    return None


def parse_sources_xlsx(content: bytes) -> tuple[list[ImportRow], list[BadRow]]:
    """Строки файла → (годные к созданию, сразу негодные). Шапка (если первая строка — не ссылка)
    и пустые строки пропускаются; нумерация — как в Excel."""
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as e:
        raise SearchInvalidParams("xlsx", "файл не открывается как Excel (.xlsx)") from e
    sheet = workbook.worksheets[0] if workbook.worksheets else None
    if sheet is None:
        raise SearchInvalidParams("xlsx", "в файле нет листов")

    rows: list[ImportRow] = []
    bad: list[BadRow] = []
    first_filled = True
    for number, cells in enumerate(sheet.iter_rows(max_col=3, values_only=True), start=1):
        if number > MAX_IMPORT_ROWS + 1:
            raise SearchInvalidParams("xlsx", f"не больше {MAX_IMPORT_ROWS} строк за раз")
        url, name, regional = (_text(c) for c in (*cells, None, None, None)[:3])
        if not (url or name or regional):
            continue
        if first_filled:
            first_filled = False
            if url and not _looks_like_link(url):
                continue  # шапка: «Ссылка | Название | Региональный»
        if not url:
            bad.append(BadRow(number, name or regional, "Не указана ссылка (колонка A)"))
            continue
        is_regional = parse_regional(regional)
        if is_regional is None:
            bad.append(BadRow(number, url, "Колонка C: ожидается «да» или «нет»"))
            continue
        rows.append(ImportRow(number, url, name or None, is_regional))
    workbook.close()
    return rows, bad
