from __future__ import annotations

import io
import uuid
from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, datetime

from search_service.applications.handlers.docx_writer import DocxBuilder, run
from search_service.applications.handlers.errors import domain_errors
from search_service.applications.handlers.results import (
    ResultFilter,
    ResultsHandler,
    filter_from_query,
)
from search_service.applications.handlers.save import make_title
from search_service.applications.ports.uow import UnitOfWork
from search_service.applications.queries.results import ExportSearchQuery
from search_service.applications.results.read_models import ExportFile
from search_service.domain.entities import Publication, Search, SearchResult, Source
from search_service.domain.enums import ExportFormat, Relevance, SearchMode
from search_service.domain.errors import NotFoundError
from search_service.domain.value_objects import KRASNOYARSK_TZ

RELEVANCE_RU = {
    Relevance.DIRECT: "Прямая",
    Relevance.INDIRECT: "Косвенная",
    Relevance.DISPUTED: "Спорная",
    Relevance.IRRELEVANT: "Не связана",
}

# В файле — слова, а не коды API (B2-08). Неизвестный код выводим как есть, а не теряем.
PLATFORM_RU = {
    "vk": "ВКонтакте",
    "ok": "Одноклассники",
    "max": "MAX",
    "news": "Новостные сайты",
}
TARGET_STATUS_RU = {
    "pending": "Ожидает",
    "running": "Идёт поиск",
    "done": "Готово",
    "empty": "Публикаций за период нет",
    "failed": "Не удалось",
    "skipped": "Пропущено",
}
SENTIMENT_RU = {"positive": "позитивная", "neutral": "нейтральная", "negative": "негативная"}
MODE_RU = {
    SearchMode.SELECTED: "по выбранным источникам",
    SearchMode.GLOBAL: "без ограничения списком",
}

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
LOCAL_TZ_NOTE = "время Красноярска"


def platform_ru(code: str) -> str:
    return PLATFORM_RU.get(code, code)


def local_dt(moment: datetime) -> str:
    """Дата и время публикации по Красноярску: «19.10.2026 21:32»."""
    return moment.astimezone(KRASNOYARSK_TZ).strftime("%d.%m.%Y %H:%M")


@dataclass
class ExportRow:
    """Одна строка подборки — ровно то, что видно на экране (тот же ListResults)."""

    number: int
    result: SearchResult
    publication: Publication
    source_name: str
    source_url: str
    title: str
    excerpt: str
    published: str  # уже по Красноярску или «недоступна»
    relevance: str
    region: str
    sentiment: str
    story: str  # номер сюжета среди выгруженных строк ("" — без группы)


@dataclass
class ExportData:
    search: Search
    rows: list[ExportRow]
    counts: dict[str, int]
    generated_at: str

    def header_lines(self) -> list[tuple[str, str]]:
        search, counts = self.search, self.counts
        return [
            ("Запрос", search.query),
            ("Как понят запрос", (search.query_analysis or {}).get("summary", "")),
            (
                "Период",
                f"{search.period.from_date.strftime('%d.%m.%Y')} — "
                f"{search.period.to_date.strftime('%d.%m.%Y')} ({LOCAL_TZ_NOTE})",
            ),
            ("Режим", MODE_RU.get(search.mode, str(search.mode))),
            ("Площадки", ", ".join(platform_ru(p) for p in search.platforms)),
            ("Сформировано", f"{self.generated_at} ({LOCAL_TZ_NOTE})"),
            (
                "Публикаций в подборке",
                f"{counts['total']} (из них без даты: {counts['unknown_date']}; "
                f"исключено вручную: {counts['excluded']})",
            ),
        ]


class ExportHandler:
    def __init__(
        self, uow: UnitOfWork, now: Callable[[], datetime] = lambda: datetime.now(UTC)
    ) -> None:
        self.uow = uow
        self.results = ResultsHandler(uow)
        self._now = now

    async def load(
        self, search_id: uuid.UUID, f: ResultFilter, sort: str = "relevance"
    ) -> ExportData:
        """Строки — тем же запросом, что и для экрана (без лимита страницы)."""
        results, _, _, counts = await self.results.list_results(
            search_id, f, sort, limit=10000, offset=0, max_limit=10000
        )
        async with self.uow:
            search = await self.uow.searches.get_with_targets(search_id)
            if not search:
                raise NotFoundError("SEARCH_NOT_FOUND")
            pubs = await self.uow.publications.get_by_ids([r.publication_id for r in results])
            pubs_by_id = {p.id: p for p in pubs}
            sources: dict[tuple[str, str], Source] = {}
            for key in {(p.platform, p.source_external_id) for p in pubs}:
                source = await self.uow.sources.get_by_key(*key)
                if source:
                    sources[key] = source

        stories: dict[str, int] = {}
        rows: list[ExportRow] = []
        for result in results:
            pub = pubs_by_id.get(result.publication_id)
            if not pub:
                continue
            source = sources.get((pub.platform, pub.source_external_id))
            if result.excerpt_start is not None and result.excerpt_end is not None:
                excerpt = pub.text[result.excerpt_start : result.excerpt_end]
            else:
                excerpt = pub.text[:300]
            region = ""
            if result.region_by_source or result.region_by_content:
                region = "Красноярский край"
                if result.region_evidence:
                    region += f" ({result.region_evidence})"
            story = ""
            if result.group_id:
                story = str(stories.setdefault(result.group_id, len(stories) + 1))
            sentiment = str(result.sentiment) if result.sentiment else ""
            rows.append(
                ExportRow(
                    number=len(rows) + 1,
                    result=result,
                    publication=pub,
                    source_name=source.name if source else pub.source_name,
                    source_url=source.url if source else pub.source_url,
                    title=make_title(pub)[0],
                    excerpt=excerpt,
                    published=local_dt(pub.published_at) if pub.published_at else "недоступна",
                    relevance=RELEVANCE_RU.get(result.relevance, "") if result.relevance else "",
                    region=region,
                    sentiment=SENTIMENT_RU.get(sentiment, sentiment),
                    story=story,
                )
            )
        generated = self._now().astimezone(KRASNOYARSK_TZ).strftime("%d.%m.%Y %H:%M")
        return ExportData(search=search, rows=rows, counts=counts, generated_at=generated)

    @staticmethod
    def filename(search: Search, extension: str) -> str:
        return f"podborka_{search.period.to_date.isoformat()}_{str(search.id)[:8]}.{extension}"

    async def to_xlsx(
        self, search_id: uuid.UUID, f: ResultFilter, sort: str = "relevance"
    ) -> tuple[bytes, str, str]:
        data = await self.load(search_id, f, sort)
        search = data.search

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()

        # ---- Лист «Подборка» ----
        ws = wb.active
        assert ws is not None
        ws.title = "Подборка"
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

        header_lines = data.header_lines()
        row = 1
        for label, value in header_lines:
            ws.cell(row=row, column=1, value=label).font = header_font
            ws.cell(row=row, column=2, value=value)
            row += 1
        row += 1

        columns = [
            "№",
            "Площадка",
            "Источник",
            "Ссылка на источник",
            "Заголовок / начало поста",
            f"Дата публикации ({LOCAL_TZ_NOTE})",
            "Ссылка на публикацию",
            "Выдержка",
            "Связь с запросом",
            "Почему",
            "Регион",
            "Тональность",
            "Сюжет №",
        ]
        for col_idx, name in enumerate(columns, 1):
            cell = ws.cell(row=row, column=col_idx, value=name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

        for r in data.rows:
            values = [
                r.number,
                platform_ru(r.publication.platform),
                r.source_name,
                r.source_url,
                r.title,
                r.published,
                r.publication.url,
                r.excerpt,
                r.relevance,
                r.result.explanation,
                r.region,
                r.sentiment,
                r.story,
            ]
            for col_idx, cell_value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=cell_value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if col_idx == 7 and cell_value:
                    cell.hyperlink = str(cell_value)
                    cell.font = Font(color="0563C1", underline="single")
            row += 1

        widths = [5, 14, 20, 30, 40, 18, 40, 50, 14, 40, 20, 14, 10]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        header_row = 1 + len(header_lines) + 1
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{row - 1}"
        ws.freeze_panes = f"A{header_row + 1}"

        # ---- Лист «Охват» ----
        ws2 = wb.create_sheet("Охват")
        cov_cols = ["Площадка", "Источник", "Статус", "Получено", "В подборке", "Ошибка"]
        for i, name in enumerate(cov_cols, 1):
            cell = ws2.cell(row=1, column=i, value=name)
            cell.font = header_font
            cell.fill = header_fill
        r2 = 2
        for t in search.targets:
            ws2.cell(row=r2, column=1, value=platform_ru(t.platform))
            ws2.cell(row=r2, column=2, value=t.name)
            ws2.cell(row=r2, column=3, value=TARGET_STATUS_RU.get(t.status.value, t.status.value))
            ws2.cell(row=r2, column=4, value=t.fetched)
            ws2.cell(row=r2, column=5, value=t.matched)
            ws2.cell(row=r2, column=6, value=t.error_message or "")
            r2 += 1
        for i, w in enumerate([16, 30, 24, 10, 12, 50], 1):
            ws2.column_dimensions[get_column_letter(i)].width = w

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), self.filename(search, "xlsx"), XLSX_MIME

    async def to_docx(
        self, search_id: uuid.UUID, f: ResultFilter, sort: str = "relevance"
    ) -> tuple[bytes, str, str]:
        """Шапка как в XLSX + список публикаций: заголовок-ссылка, источник, дата, выдержка."""
        data = await self.load(search_id, f, sort)
        doc = DocxBuilder()
        doc.title("Подборка публикаций")
        for label, value in data.header_lines():
            doc.paragraph(run(f"{label}: ", bold=True), value)
        doc.heading("Публикации")
        if not data.rows:
            doc.paragraph("Подходящих публикаций не найдено.")
        for r in data.rows:
            doc.paragraph(
                run(f"{r.number}. ", bold=True),
                run(r.title, bold=True, link=r.publication.url),
            )
            details = [platform_ru(r.publication.platform), r.source_name, r.published]
            if r.relevance:
                details.append(f"связь с запросом: {r.relevance.lower()}")
            doc.paragraph(run(" · ".join(details), italic=True))
            doc.paragraph(r.excerpt)
            if r.result.explanation:
                doc.paragraph(run("Почему: ", bold=True), r.result.explanation)
        return doc.build(), self.filename(data.search, "docx"), DOCX_MIME


# ---------------------------------------------------------------------------
# Хэндлер медиатора для gRPC-транспорта (регистрируется в dependencies.build_mediator).
# ---------------------------------------------------------------------------


@dataclass
class ExportSearchQueryHandler:
    uow: UnitOfWork

    async def handle(self, query: ExportSearchQuery) -> ExportFile:
        """Строки — тем же запросом, что ListResults (без limit), тот же filter и sort →
        файл совпадает с экраном. Формат листов — B2-08."""
        handler = ExportHandler(self.uow)
        render = handler.to_docx if query.format is ExportFormat.DOCX else handler.to_xlsx
        with domain_errors():
            content, filename, mime = await render(
                query.search_id, filter_from_query(query.filter), sort=query.sort.value
            )
        return ExportFile(content=content, filename=filename, mime_type=mime)
