"""XLSX (B2-08): строки и порядок ссылок = ListResults, слова вместо кодов, время Красноярска."""

import io
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from search_service.applications.handlers.export import ExportHandler
from search_service.applications.handlers.results import ResultFilter, ResultsHandler
from search_service.applications.handlers.scoring import ScoringHandler
from search_service.applications.ports.ml_client import QueryAnalysis
from search_service.domain.enums import QueryType, Relevance
from tests.builders import add_result, add_search, add_source
from tests.fakes import FakeMl, InMemoryUoW

NOW = datetime(2026, 10, 19, 14, 0, tzinfo=UTC)  # 21:00 в Красноярске
ANALYSIS = QueryAnalysis(QueryType.TOPIC.value, "Безопасность на дорогах", ["дороги"])


@pytest.fixture
async def stand():
    """Поиск с двумя целями (одна упала) и тремя оценёнными постами разной связи."""
    uow, ml = InMemoryUoW(), FakeMl()
    good, bad = await add_source(uow, "good"), await add_source(uow, "bad")
    search = await add_search(uow, [good, bad])
    ok, failed = search.targets
    ok.mark_done(fetched=3)
    failed.mark_failed("RATE_LIMITED", "Превышен лимит запросов к площадке")
    results = [
        await add_result(uow, search, ok, n, f"Пост {n} про дороги", source="good")
        for n in (1, 2, 3)
    ]
    ml.relevance_by_text = {"Пост 2": Relevance.INDIRECT, "Пост 3": Relevance.DISPUTED}
    await ScoringHandler(uow, ml).score_search(search.id, ANALYSIS)
    results[0].sentiment = "positive"
    results[0].group_id = "g-a"
    results[1].group_id = "g-a"
    results[2].group_id = "g-b"
    return uow, search, results


def sheet(content: bytes, name: str):
    return load_workbook(io.BytesIO(content))[name]


async def test_xlsx_rows_and_links_match_list_results(stand):
    uow, search, _ = stand
    content, filename, mime = await ExportHandler(uow, now=lambda: NOW).to_xlsx(
        search.id, ResultFilter()
    )

    shown, total, _, counts = await ResultsHandler(uow).list_results(search.id, ResultFilter())
    ws = sheet(content, "Подборка")
    links = [c.value for c in ws["G"] if isinstance(c.value, str) and c.value.startswith("http")]

    assert len(links) == counts["total"] == total == 3
    assert links == [f"https://vk.com/wallgood_{n}" for n in (1, 2, 3)]
    assert links == [(await uow.publications.get_by_ids([r.publication_id]))[0].url for r in shown]
    assert [c.hyperlink.target for c in ws["G"] if c.hyperlink] == links
    assert filename.startswith("podborka_") and filename.endswith(".xlsx") and filename.isascii()
    assert "spreadsheetml" in mime


async def test_xlsx_has_words_not_codes(stand):
    uow, search, _ = stand
    content, _, _ = await ExportHandler(uow, now=lambda: NOW).to_xlsx(search.id, ResultFilter())

    coverage = [[c.value for c in row] for row in sheet(content, "Охват").iter_rows(min_row=2)]
    assert [row[0] for row in coverage] == ["ВКонтакте", "ВКонтакте"]
    assert [row[2] for row in coverage] == ["Готово", "Не удалось"]
    assert coverage[1][5] == "Превышен лимит запросов к площадке"

    selection = [[c.value for c in row] for row in sheet(content, "Подборка").iter_rows()]
    flat = {str(v) for row in selection for v in row if v is not None}
    assert "ВКонтакте" in flat
    assert {"Прямая", "Косвенная", "Спорная", "позитивная"} <= flat
    assert not flat & {"vk", "VK", "done", "failed", "positive", "direct"}


async def test_xlsx_header_uses_krasnoyarsk_time_and_story_numbers(stand):
    uow, search, _ = stand
    content, _, _ = await ExportHandler(uow, now=lambda: NOW).to_xlsx(search.id, ResultFilter())
    ws = sheet(content, "Подборка")

    header = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value for i in range(1, 8)}
    assert header["Сформировано"] == "19.10.2026 21:00 (время Красноярска)"
    assert header["Площадки"] == "ВКонтакте"
    assert header["Публикаций в подборке"] == "3 (из них без даты: 0; исключено вручную: 0)"
    # сюжеты нумеруются по порядку строк: два поста в g-a — сюжет 1, третий — сюжет 2
    stories = [c.value for c in ws["M"][9:]]
    assert stories == ["1", "1", "2"]


async def test_xlsx_publication_date_is_local_time(stand):
    uow, search, results = stand
    (pub,) = await uow.publications.get_by_ids([results[0].publication_id])
    pub.published_at = datetime(2026, 10, 19, 16, 30, tzinfo=UTC)
    pub2 = (await uow.publications.get_by_ids([results[1].publication_id]))[0]
    pub2.published_at = None

    content, _, _ = await ExportHandler(uow, now=lambda: NOW).to_xlsx(search.id, ResultFilter())
    dates = [c.value for c in sheet(content, "Подборка")["F"][9:]]

    assert dates[0] == "19.10.2026 23:30"
    assert dates[1] == "недоступна"


async def test_docx_is_valid_and_lists_publications_with_links(stand):
    """DOCX (B2-10): шапка как в XLSX + публикации; ссылки — настоящие гиперссылки."""
    import zipfile
    from xml.etree import ElementTree

    uow, search, _ = stand
    content, filename, mime = await ExportHandler(uow, now=lambda: NOW).to_docx(
        search.id, ResultFilter()
    )

    assert filename.endswith(".docx") and filename.isascii() and "wordprocessingml" in mime
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    for xml in parts.values():
        ElementTree.fromstring(xml)  # каждая часть — well-formed XML
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    body = ElementTree.fromstring(parts["word/document.xml"])
    text = "\n".join(
        "".join(t.text or "" for t in p.iter(f"{{{ns['w']}}}t"))
        for p in body.iter(f"{{{ns['w']}}}p")
    )
    assert "Сформировано: 19.10.2026 21:00 (время Красноярска)" in text
    assert "ВКонтакте" in text and "Прямая".lower() in text.lower()
    rels = ElementTree.fromstring(parts["word/_rels/document.xml.rels"])
    targets = [r.get("Target") for r in rels if r.get("Type", "").endswith("/hyperlink")]
    assert targets == [f"https://vk.com/wallgood_{n}" for n in (1, 2, 3)]
